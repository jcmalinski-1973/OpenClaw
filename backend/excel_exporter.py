import openpyxl


def export_adjustment_excel(template_path: str, output_path: str, results: list[dict]) -> None:
    workbook = openpyxl.load_workbook(template_path)
    sheet = workbook[workbook.sheetnames[0]]

    for row_idx in range(4, sheet.max_row + 1):
        sheet.cell(row_idx, 4).value = None

    for item in results:
        if item.get("status") != "ok":
            continue
        sheet.cell(item["template_row"], 4).value = item["additional_packs"]

    workbook.save(output_path)
