import math
import re
import unicodedata
import difflib
import openpyxl


STOPWORDS = {
    "BREAD",
    "KING",
    "BK",
    "UN",
    "UNID",
    "UNID.",
    "PCT",
    "PCTS",
    "CX",
    "FD",
    "KG",
    "G",
    "C",
    "C/",
    "CXS",
    "FARDO",
    "PACOTE",
    "PCTE",
    "PRE",
    "ASSADO",
    "ULTRACONGELADO",
    "CONG",
    "CONGELADO",
    "TRAD",
    "TRADICIONAL",
}


def normalize_code(value) -> str:
    text = str(value).strip()
    return text.lstrip("0") or "0"


def normalize_description(text: str) -> str:
    text = unicodedata.normalize("NFKD", text.upper()).encode("ASCII", "ignore").decode("ASCII")
    text = text.replace("REQUEIJAO", "REQ").replace("CATUPIRY", "REQ")
    text = re.sub(r"\b\d+[.,]?\d*[A-Z/]*\b", " ", text)
    text = re.sub(r"[^A-Z0-9 ]+", " ", text)
    tokens = [token for token in text.split() if len(token) > 1 and token not in STOPWORDS]
    return " ".join(tokens)


def load_template_rows(template_path: str):
    workbook = openpyxl.load_workbook(template_path)
    sheet = workbook[workbook.sheetnames[0]]

    rows = {}
    for row_idx in range(4, sheet.max_row + 1):
        code = sheet.cell(row_idx, 1).value
        if code is None:
            continue

        normalized_code = normalize_code(code)

        rows[normalized_code] = {
            "row": row_idx,
            "code_raw": str(code).strip(),
            "description": sheet.cell(row_idx, 2).value or "",
            "pack_size": int(sheet.cell(row_idx, 3).value or 0),
        }

    return workbook, sheet, rows


def best_stock_match(order_item: dict, stock_items: list[dict]):
    exact = next((item for item in stock_items if item["code"] == order_item["code"]), None)

    if exact and (
        exact["mes"] > 0
        or exact["m1"] > 0
        or exact["m2"] > 0
        or exact["m3"] > 0
        or exact["disp"] > 0
    ):
        return exact, "exact"

    order_norm = normalize_description(order_item["description"])
    order_tokens = set(order_norm.split())

    best_item = None
    best_score = 0.0

    for stock_item in stock_items:
        stock_norm = normalize_description(stock_item["description"])
        stock_tokens = set(stock_norm.split())

        if not stock_tokens:
            continue

        overlap = order_tokens & stock_tokens
        if len(overlap) < 2:
            continue

        sequence_score = difflib.SequenceMatcher(None, order_norm, stock_norm).ratio()
        overlap_score = len(overlap) / max(len(order_tokens), 1)
        score = (sequence_score * 0.6) + (overlap_score * 0.4)

        if stock_item["mes"] or stock_item["m1"] or stock_item["m2"] or stock_item["m3"]:
            score += 0.05

        if score > best_score:
            best_score = score
            best_item = stock_item

    if best_item and best_score >= 0.55:
        return best_item, f"desc:{best_score:.2f}"

    if exact:
        return exact, "exact-zero"

    return None, "none"


def compute_adjustments(order_items: list[dict], stock_data: dict, template_rows: dict) -> list[dict]:
    report_date = stock_data.get("report_date")
    report_day = report_date.day if report_date else 30
    results = []

    for order_item in order_items:
        template = template_rows.get(order_item["code"])

        if not template:
            results.append(
                {
                    **order_item,
                    "status": "template_not_found",
                }
            )
            continue

        stock_item, match_type = best_stock_match(order_item, stock_data["items"])

        if not stock_item:
            results.append(
                {
                    **order_item,
                    "template_row": template["row"],
                    "status": "stock_not_found",
                }
            )
            continue

        pack_size = template["pack_size"] or 1
        suggested_units = order_item["suggested_packs"] * pack_size

        average_last_3_months = (stock_item["m1"] + stock_item["m2"] + stock_item["m3"]) / 3
        daily_historical = average_last_3_months / 30
        daily_current = stock_item["mes"] / max(report_day, 1)
        daily_final = max(daily_historical, daily_current)

        target_units_20_days = math.ceil(daily_final * 20)
        additional_units = max(target_units_20_days - stock_item["disp"] - suggested_units, 0)
        additional_packs = math.ceil(additional_units / pack_size) if additional_units > 0 else 0

        results.append(
            {
                **order_item,
                "template_row": template["row"],
                "template_code": template["code_raw"],
                "template_description": template["description"],
                "pack_size": pack_size,
                "stock_code": stock_item["code"],
                "stock_description": stock_item["description"],
                "disp": stock_item["disp"],
                "mes": stock_item["mes"],
                "m1": stock_item["m1"],
                "m2": stock_item["m2"],
                "m3": stock_item["m3"],
                "average_last_3_months": round(average_last_3_months, 2),
                "daily_historical": round(daily_historical, 4),
                "daily_current": round(daily_current, 4),
                "daily_final": round(daily_final, 4),
                "target_units_20_days": target_units_20_days,
                "suggested_units": suggested_units,
                "additional_units": additional_units,
                "additional_packs": additional_packs,
                "match_type": match_type,
                "status": "ok",
            }
        )

    return results
